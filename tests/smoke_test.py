#!/usr/bin/env python3
"""Standalone clean-room smoke test for the published skill repository."""
from __future__ import annotations

import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET

import docx
import openpyxl


ROOT = Path(__file__).resolve().parents[1]
MODEL = ROOT / "examples" / "corporate-purchase-planning-model.json"


def run(*args: str, cwd: Path, env: dict | None = None) -> subprocess.CompletedProcess:
    merged = os.environ.copy()
    merged.update({"PYTHONUTF8": "1", "PYTHONIOENCODING": "utf-8"})
    if env:
        merged.update(env)
    result = subprocess.run([sys.executable, *args], cwd=cwd, env=merged,
                            capture_output=True, text=True, encoding="utf-8", errors="replace")
    if result.returncode:
        raise RuntimeError(result.stdout + result.stderr)
    return result


def main() -> int:
    with tempfile.TemporaryDirectory(prefix="bpmn-skill-smoke-") as raw:
        workspace = Path(raw)
        text_source = workspace / "source.txt"
        text_source.write_text("Получить заказ\nОтгрузить товар", encoding="utf-8")
        document_source = workspace / "source.docx"
        document = docx.Document()
        document.add_heading("Процесс", level=1)
        document.add_paragraph("Создать заказ")
        document.save(document_source)
        workbook_source = workspace / "source.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.title = "План"
        workbook.active["A1"] = "Шаг"
        workbook.active["B2"] = "Создать заказ"
        workbook.save(workbook_source)
        extractor = ROOT / "scripts" / "extract_protocol.py"
        for source in (text_source, document_source, workbook_source):
            extracted = json.loads(run(str(extractor), str(source), cwd=workspace).stdout)
            if not extracted.get("raw_text"):
                raise RuntimeError(f"Extraction produced no text for {source.name}")
        spreadsheet = json.loads(run(str(extractor), str(workbook_source), cwd=workspace).stdout)
        if "'План'!B2" not in {item["locator"] for item in spreadsheet["locators"]}:
            raise RuntimeError("Spreadsheet cell locator is missing")
        model = workspace / "model.json"
        shutil.copy2(MODEL, model)
        run(str(ROOT / "scripts" / "run_pipeline.py"), "--model", str(model),
            "--workspace", str(workspace), "--no-png", "--bpmn", cwd=workspace,
            env={"BPMN_DISABLE_DRAWIO": "1"})
        process_id = "corporate-purchase-planning"
        drawio = workspace / "output" / "drawio" / f"{process_id}.drawio"
        bpmn = workspace / "output" / "bpmn" / f"{process_id}.bpmn"
        report_path = workspace / "output" / "validation" / f"{process_id}-drawio-validation.json"
        if not drawio.is_file() or not bpmn.is_file():
            raise RuntimeError("Expected draw.io/BPMN deliverables are missing")
        report = json.loads(report_path.read_text(encoding="utf-8"))
        if report.get("status") != "NEEDS_REVIEW":
            raise RuntimeError("Missing draw.io CLI must produce NEEDS_REVIEW")
        root = ET.parse(bpmn).getroot()
        if root.find("{http://www.omg.org/spec/BPMN/20100524/DI}BPMNDiagram") is None:
            raise RuntimeError("BPMN-DI is missing")
    print("Standalone smoke test: PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
